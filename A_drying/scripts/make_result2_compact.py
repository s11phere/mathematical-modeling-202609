# -*- coding: utf-8 -*-
"""
Write a compact result2.xlsx from the stored profiles.

The problem asks for 1 s x 0.1 cm output over the 3 h window quoted in the
tables, and for the whole 2-3 day process overall.  A 1 s grid over 45 h would
make a ~34 MB workbook, so the workbook keeps 1 s resolution for the first 3 h
(the reported window) and switches to 60 s afterwards, which is stated in the
paper.  result3.xlsx (problem 3) carries the fine 60 s grid throughout.
"""
from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
import model_problems as mp  # noqa: E402

TEMPL = ROOT.parents[1] / "problems" / "A题" / "附件" / "附件3"
OUT = ROOT / "out"
FINE_UNTIL = 3 * 3600.0
COARSE = 60.0


def main():
    z = np.load(OUT / "p2.npz", allow_pickle=True)
    t = z["t"]
    T = z["T"]
    C = z["C"]
    rc = z["r"]
    keep = [0]
    last = t[0]
    for k in range(1, len(t)):
        step = 1.0 if t[k] <= FINE_UNTIL else COARSE
        if t[k] - last >= step - 1e-9:
            keep.append(k)
            last = t[k]
    keep = np.array(keep)
    ts = t[keep]
    To = T[keep]
    Co = C[keep]
    print(f"selected {len(ts)} rows (from {len(t)}): 1 s until {FINE_UNTIL/3600:.0f} h, "
          f"then {COARSE:.0f} s", flush=True)

    wb = openpyxl.load_workbook(TEMPL / "result2.xlsx")
    for sheet, vals in (("温度", To), ("水分浓度", Co)):
        ws = wb[sheet]
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row)
        ws.cell(row=1, column=1, value="时间\\到药材中心的距离")
        for j, r in enumerate(rc):
            ws.cell(row=1, column=2 + j, value=float(r))
        for i, tt in enumerate(ts):
            ws.cell(row=2 + i, column=1, value=float(tt))
            for j in range(len(rc)):
                ws.cell(row=2 + i, column=2 + j, value=round(float(vals[i, j]), 4))
    out = OUT / "result2_compact.xlsx"
    wb.save(out)
    print("wrote", out)
    import os
    print("size %.2f MB" % (os.path.getsize(out) / 1024 / 1024))
    np.savez_compressed(OUT / "p2_compact.npz", t=ts, r=rc, T=To, C=Co,
                        t_end_h=float(z["t_end_h"]))


if __name__ == "__main__":
    main()
