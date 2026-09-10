# -*- coding: utf-8 -*-
"""
Write the answer workbooks in the format of 附件 3 (result1..result4.xlsx).

Templates (read-only input):
    problems/A题/附件/附件3/result1.xlsx   sheets 温度 / 水分浓度
    problems/A题/附件/附件3/result2.xlsx   sheets 温度 / 水分浓度
    problems/A题/附件/附件3/result3.xlsx   sheet  Sheet1
    problems/A题/附件/附件3/result4.xlsx   sheet  Sheet1, last column 药材表面

Every template has column A = time and row 1 = distance from the centre; the
workbooks are rewritten with the full grids required by the problem statement
(1 s x 0.1 cm for problems 1 and 2; 60 s x 0.1 cm for problems 3 and 4).
All values are rounded to four decimal places, as required.
"""
from __future__ import annotations

from pathlib import Path
from typing import Sequence

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[3]     # -> project root
TEMPLATES = ROOT / "problems" / "A题" / "附件" / "附件3"


def _clean(ws, n_row: int, n_col: int) -> None:
    """Delete every existing data row below the header."""
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row)


def _fill(ws, times, r_cm, values, surface_label: str | None = None) -> None:
    """Write the header and the table into ``ws``."""
    ws.cell(row=1, column=1, value="时间\\到药材中心的距离")
    for j, r in enumerate(r_cm):
        ws.cell(row=1, column=2 + j, value=float(r))
    if surface_label is not None:
        ws.cell(row=1, column=2 + len(r_cm), value=surface_label)
    for i, t in enumerate(times):
        ws.cell(row=2 + i, column=1, value=float(t))
        for j in range(len(r_cm)):
            ws.cell(row=2 + i, column=2 + j, value=round(float(values[i, j]), 4))


def write_result1(path: Path, times, r_cm, T_C, C):
    wb = openpyxl.load_workbook(TEMPLATES / "result1.xlsx")
    for sheet, vals in (("温度", T_C), ("水分浓度", C)):
        ws = wb[sheet] if sheet in wb.sheetnames else wb.create_sheet(sheet)
        _clean(ws, len(times), len(r_cm))
        _fill(ws, times, r_cm, vals)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_result2(path: Path, times, r_cm, T_C, C):
    wb = openpyxl.load_workbook(TEMPLATES / "result2.xlsx")
    for sheet, vals in (("温度", T_C), ("水分浓度", C)):
        ws = wb[sheet] if sheet in wb.sheetnames else wb.create_sheet(sheet)
        _clean(ws, len(times), len(r_cm))
        _fill(ws, times, r_cm, vals)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_result3(path: Path, times, r_cm, C):
    wb = openpyxl.load_workbook(TEMPLATES / "result3.xlsx")
    ws = wb.worksheets[0]
    _clean(ws, len(times), len(r_cm))
    _fill(ws, times, r_cm, C)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_result4(path: Path, times, r_cm, C, C_surface):
    """Problem 4: the last column is the moving surface value ``药材表面``."""
    wb = openpyxl.load_workbook(TEMPLATES / "result4.xlsx")
    ws = wb.worksheets[0]
    _clean(ws, len(times), len(r_cm) + 1)
    _fill(ws, times, r_cm, C, surface_label="药材表面")
    for i in range(len(times)):
        ws.cell(row=2 + i, column=2 + len(r_cm),
                value=round(float(C_surface[i]), 4))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
