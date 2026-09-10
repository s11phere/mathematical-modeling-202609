# -*- coding: utf-8 -*-
"""
Produce result4.xlsx (problem 4) and determine the true drying time.

附件2 ends at 72 h with R = 1.198 cm and dR/dt = 0 thereafter (the measured
radius is constant to 3 decimals from 40 h on), so the model is continued with
R held at its final value; this is stated in the paper.
"""
from __future__ import annotations

import json
import shutil
import sys
import time
from pathlib import Path

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from mol_solver import MolSolver, Series, c2k, k2c  # noqa: E402
import model_problems as mp  # noqa: E402

TEMPL = ROOT.parents[1] / "problems" / "A题" / "附件" / "附件3"
OUT = ROOT / "out"
T_PHASE = 1800.0
T_DATA_END = 259200.0     # 附件2 ends here (72 h)
N = 150
DAYS = 12.0


class RFun:
    """R(t) from 附件2, held constant at its final value beyond the data."""

    def __init__(self, smooth=True):
        self.Rf, self.dRf = mp.make_Rfun(smooth=smooth)

    def __call__(self, t):
        return float(self.Rf(min(float(t), T_DATA_END)))


def fill(ws, times, r_cm, vals, surf=None):
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row)
    ws.cell(row=1, column=1, value="时间\\到药材中心的距离")
    for j, r in enumerate(r_cm):
        ws.cell(row=1, column=2 + j, value=float(r))
    if surf is not None:
        ws.cell(row=1, column=2 + len(r_cm), value="药材表面")
    for i, t in enumerate(times):
        ws.cell(row=2 + i, column=1, value=float(t))
        for j in range(len(r_cm)):
            ws.cell(row=2 + i, column=2 + j, value=round(float(vals[i, j]), 4))
    if surf is not None:
        for i in range(len(times)):
            ws.cell(row=2 + i, column=2 + len(r_cm),
                    value=round(float(surf[i]), 4))


def main():
    amb = mp.Ambient()
    Rfun = RFun()
    R_end = float(Rfun(T_DATA_END))
    rc = mp.output_radii_cm(0.1, round(R_end * 100, 1))
    r_out = np.asarray(rc) * 1e-2
    print(f"R_end = {R_end*100:.3f} cm ; grid 0..{rc[-1]} cm ({len(rc)} cols)",
          flush=True)

    pa = mp.props_problem1(amb)
    pa.hm_mode = "heat_limited"
    sa = MolSolver(N=N, R0=mp.R0, Rfun=Rfun)
    t0 = time.time()
    ra = sa.run(pa, T_PHASE, c2k(mp.T_INIT_C), mp.C_INIT, t_start=0.0,
                sample_dt=60.0, rtol=1e-7, atol=1e-9, method="BDF", max_step=300.0)
    pb = mp.props_appendix4(amb)
    pb.hm_mode = "heat_limited"
    sb = MolSolver(N=N, R0=mp.R0, Rfun=Rfun)
    rb = sb.run(pb, DAYS * 24 * 3600.0, ra.T[-1], ra.C[-1], t_start=T_PHASE,
                sample_dt=60.0, rtol=1e-6, atol=1e-8, method="BDF", max_step=900.0)
    print(f"integration took {time.time()-t0:.1f} s", flush=True)

    ts = list(ra.ts) + list(rb.ts)[1:]
    C = list(ra.C) + list(rb.C)[1:]
    Cs = list(ra.Csurf) + list(rb.Csurf)[1:]
    Rs = list(ra.R) + list(rb.R)[1:]
    nT = len(ts)
    Co = np.empty((nT, len(r_out)))
    for k in range(nT):
        r_nodes = np.concatenate(([0.0], Rs[k] * sa.xi_c, [Rs[k]]))
        Cn = np.concatenate(([C[k][0]], C[k][:N], [Cs[k]]))
        Co[k] = np.interp(r_out, r_nodes, Cn)
    Cmax = np.maximum(Co.max(axis=1), np.array(Cs))
    hit = np.where(Cmax < mp.C_TARGET)[0]
    t_end_h = float(ts[hit[0]] / 3600.0) if hit.size else None

    wb = openpyxl.load_workbook(TEMPL / "result4.xlsx")
    fill(wb.worksheets[0], ts, rc, Co, surf=Cs)
    tmp = OUT / "result4_tmp.xlsx"
    wb.save(tmp)
    np.savez_compressed(OUT / "p4.npz", t=np.array(ts), r=np.array(rc), C=Co,
                        Cs=np.array(Cs), R=np.array(Rs),
                        t_end_h=(t_end_h if t_end_h else np.nan))
    try:
        shutil.copyfile(tmp, OUT / "result4.xlsx")
        print("copied to result4.xlsx", flush=True)
    except PermissionError:
        print("result4.xlsx locked; left as result4_tmp.xlsx", flush=True)

    print(f"t_end(0.15) = {t_end_h} h ; Cmax_end = {Cmax[-1]:.4f}", flush=True)
    for h in (6, 12, 24, 36, 48, 60, 72, 96, 120, 144):
        if h * 3600.0 <= ts[-1]:
            k = int(np.argmin(np.abs(np.array(ts) - h * 3600.0)))
            print(f"  {h:4d} h  core={Co[k][0]:.4f}  surf={Cs[k]:.4f}  "
                  f"R={Rs[k]*100:.3f} cm", flush=True)
    (OUT / "p4_meta.json").write_text(
        json.dumps(dict(N=N, n=nT, t_end_h=t_end_h,
                        Cmax_end=float(Cmax[-1]),
                        R_end_cm=float(Rs[-1] * 100),
                        wall=time.time() - t0),
                   indent=2, ensure_ascii=False), encoding="utf-8")


if __name__ == "__main__":
    main()
