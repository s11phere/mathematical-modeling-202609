# -*- coding: utf-8 -*-
"""
Produce result3.xlsx (problem 3) writing first to a temporaty path so that a
lock on the final file (e.g. an open Excel window) does not lose the run.
"""
from __future__ import annotations

import json
import shutil
import sys
import time
from pathlib import Path

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from mol_solver import MolSolver, Series, c2k, k2c  # noqa: E402
import model_problems as mp  # noqa: E402

TEMPL = ROOT.parents[1] / "problems" / "A题" / "附件" / "附件3"
OUT = ROOT / "out"
T_PHASE = 1800.0
N = 150


def fill(ws, times, r_cm, vals):
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row)
    ws.cell(row=1, column=1, value="时间\\到药材中心的距离")
    for j, r in enumerate(r_cm):
        ws.cell(row=1, column=2 + j, value=float(r))
    for i, t in enumerate(times):
        ws.cell(row=2 + i, column=1, value=float(t))
        for j in range(len(r_cm)):
            ws.cell(row=2 + i, column=2 + j, value=round(float(vals[i, j]), 4))


def main():
    amb = mp.Ambient()
    rc = mp.output_radii_cm(0.1, 2.0)
    r_out = np.asarray(rc) * 1e-2

    pa = mp.props_problem1(amb)
    pa.hm_mode = "heat_limited"
    sa = MolSolver(N=N, R0=mp.R0)
    t0 = time.time()
    ra = sa.run(pa, T_PHASE, c2k(mp.T_INIT_C), mp.C_INIT, t_start=0.0,
                sample_dt=60.0, rtol=1e-7, atol=1e-9, method="BDF", max_step=300.0)
    pb = mp.props_appendix3(amb)
    pb.hm_mode = "heat_limited"
    sb = MolSolver(N=N, R0=mp.R0)
    rb = sb.run(pb, 30 * 24 * 3600.0, ra.T[-1], ra.C[-1], t_start=T_PHASE,
                sample_dt=60.0, rtol=1e-6, atol=1e-8, method="BDF", max_step=600.0)
    print(f"integration took {time.time()-t0:.1f} s", flush=True)

    ts = list(ra.ts) + list(rb.ts)[1:]
    C = list(ra.C) + list(rb.C)[1:]
    Cs = list(ra.Csurf) + list(rb.Csurf)[1:]
    nT = len(ts)
    Co = np.empty((nT, len(r_out)))
    for k in range(nT):
        r_nodes = np.concatenate(([0.0], mp.R0 * sa.xi_c, [mp.R0]))
        Cn = np.concatenate(([C[k][0]], C[k][:N], [Cs[k]]))
        Co[k] = np.interp(r_out, r_nodes, Cn)
    Cmax = np.max(Co, axis=1)
    hit = np.where(Cmax < mp.C_TARGET)[0]
    t_end_h = float(ts[hit[0]] / 3600.0) if hit.size else None

    wb = openpyxl.load_workbook(TEMPL / "result3.xlsx")
    fill(wb.worksheets[0], ts, rc, Co)
    tmp = OUT / "result3_tmp.xlsx"
    wb.save(tmp)
    print(f"saved {tmp}  rows={len(ts)}  t_end_h={t_end_h}  "
          f"Cmax_end={Cmax[-1]:.4f}", flush=True)
    np.savez_compressed(OUT / "p3.npz", t=np.array(ts), r=np.array(rc), C=Co,
                        t_end_h=(t_end_h if t_end_h else np.nan))
    try:
        shutil.copyfile(tmp, OUT / "result3.xlsx")
        print("copied to result3.xlsx", flush=True)
    except PermissionError:
        print("result3.xlsx is locked; left as result3_tmp.xlsx", flush=True)
    (OUT / "p3_meta.json").write_text(
        json.dumps(dict(N=N, n=len(ts), t_end_h=t_end_h,
                        Cmax_end=float(Cmax[-1]), wall=time.time() - t0),
                   indent=2, ensure_ascii=False), encoding="utf-8")


if __name__ == "__main__":
    main()
